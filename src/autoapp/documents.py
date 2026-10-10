"""Loading and reading the user's documents (design decision 2).

``documents/`` is read-only for the program: nothing here writes into it.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from autoapp.config.models import DocumentEntry, Profile
from autoapp.config.workspace import Workspace


class DocumentError(Exception):
    """A problem with the documents folder or a document's content."""


@dataclass
class DocumentSet:
    """Declared documents, their extracted text (CVs and writing samples) and stray files."""

    directory: Path
    entries: list[DocumentEntry]
    texts: dict[str, str] = field(default_factory=dict)  # entry.file -> extracted text
    unmapped: list[str] = field(default_factory=list)

    def path_of(self, entry: DocumentEntry) -> Path:
        return self.directory / entry.file

    def by_role(self, role: str) -> list[DocumentEntry]:
        return [e for e in self.entries if e.role == role]

    @property
    def cv(self) -> DocumentEntry:
        return self.by_role("cv")[0]

    def cv_text(self, language: str | None = None) -> str:
        """Text of the CV in ``language`` if one is declared, else of the first CV."""
        cvs = self.by_role("cv")
        chosen = next((e for e in cvs if language and e.language == language), cvs[0])
        return self.texts[chosen.file]


def extract_text(path: Path) -> str:
    """Extract text from a PDF or DOCX file (empty string if there is none)."""
    path = Path(path)
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
        if suffix == ".docx":
            from docx import Document

            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.extend(c.text for c in row.cells)
            return "\n".join(p for p in parts if p.strip()).strip()
    except Exception as exc:  # unreadable/corrupt file
        raise DocumentError(f"cannot read {path}: {exc}") from exc
    raise DocumentError(f"unsupported document type '{path.suffix}' for {path}; use PDF or DOCX")


def _list_files(directory: Path) -> list[str]:
    return sorted(
        p.relative_to(directory).as_posix()
        for p in directory.rglob("*")
        if p.is_file() and not any(part.startswith(".") for part in p.relative_to(directory).parts)
    )


def load_documents(ws: Workspace, profile: Profile) -> DocumentSet:
    """Validate the declared documents and extract text from CVs and writing samples."""
    directory = ws.documents_dir
    if not directory.is_dir():
        raise DocumentError(f"documents folder not found: {directory}")
    entries = list(profile.documents)
    if not any(e.role == "cv" for e in entries):
        raise DocumentError("a CV is required: declare at least one document with role 'cv' in profile.yaml")
    missing = [e.file for e in entries if not (directory / e.file).is_file()]
    if missing:
        raise DocumentError(
            "declared documents not found in " + str(directory) + ": " + ", ".join(sorted(missing))
        )
    texts: dict[str, str] = {}
    for e in entries:
        if e.role not in ("cv", "writing_sample"):
            continue
        text = extract_text(directory / e.file)
        if not text and e.role == "cv":
            raise DocumentError(
                f"the CV '{e.file}' has no extractable text (image-only or scanned?); "
                "provide a text-based PDF or DOCX"
            )
        texts[e.file] = text
    declared = {e.file for e in entries}
    unmapped = [f for f in _list_files(directory) if f not in declared]
    return DocumentSet(directory=directory, entries=entries, texts=texts, unmapped=unmapped)


# ------------------------------------------------------------------ role proposal (2.6)

_TOKEN_SPLIT = re.compile(r"[^a-z0-9äöüß]+")


def _has(stem: str, tokens: list[str], substrings: tuple[str, ...] = (), words: tuple[str, ...] = ()) -> bool:
    return any(s in stem for s in substrings) or any(w in tokens for w in words)


def propose_roles(filenames: list[str]) -> list[DocumentEntry]:
    """Guess role and language for each file name (case-insensitive heuristics)."""
    out: list[DocumentEntry] = []
    for name in filenames:
        stem = Path(name).stem.lower()
        tokens = [t for t in _TOKEN_SPLIT.split(stem) if t]
        if _has(stem, tokens, ("foto", "photo")):
            role = "photo"
        elif _has(stem, tokens, ("arbeitszeugnis", "zwischenzeugnis", "referenz", "reference", "employment")):
            role = "work_reference"
        elif _has(stem, tokens, ("zeugnis", "certificate", "zertifikat")):
            role = "certificate"
        elif _has(stem, tokens, ("anschreiben", "cover")):
            role = "writing_sample"
        elif _has(stem, tokens, ("lebenslauf", "resume"), ("cv",)):
            role = "cv"
        else:
            role = "other"
        language = None
        if _has(stem, tokens, ("deutsch", "german"), ("de",)):
            language = "de"
        elif _has(stem, tokens, ("english", "englisch"), ("en",)):
            language = "en"
        if role not in ("cv", "writing_sample"):
            language = None  # certificates and the like are sent whatever the ad's language, unless you say otherwise
        out.append(DocumentEntry(file=name, role=role, language=language))
    return out
